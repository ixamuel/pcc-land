#!/usr/bin/env python3
"""Build the part-number catalog used by index.html.

Reads
  Inductor all data.xlsx        -> the 150 catalogue part numbers
  data/terminal_dims.csv        -> terminal dimensions per 3-letter suffix
  data/pn_overrides.csv         -> per part-number exceptions

Writes
  data/catalog.json             -> machine readable catalog
  index.html                    -> inlines the catalog between the markers
                                   /* BEGIN GENERATED CATALOG */ ... END

The catalog is inlined so index.html stays a single, offline, shareable file.

Usage
  .venv/bin/python tools/build_catalog.py
  .venv/bin/python tools/build_catalog.py --check   # fail if index.html is stale
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "Inductor all data.xlsx"
SHEET = "All Data"
DIMS_CSV = ROOT / "data" / "terminal_dims.csv"
OVERRIDES_CSV = ROOT / "data" / "pn_overrides.csv"
OUT_JSON = ROOT / "data" / "catalog.json"
HTML = ROOT / "index.html"

BEGIN = "/* BEGIN GENERATED CATALOG */"
END = "/* END GENERATED CATALOG */"

# Column layout of the "All Data" sheet (1-based).
COL = {
    "catalog": 2,      # B  カタログ品
    "status": 6,       # F  Status (MP / ...)
    "feature": 11,     # K  Feature
    "series": 12,      # L  Series
    "size": 13,        # M  Series (Size)
    "pn": 15,          # O  Part Number
    "body_l": 28,      # AB L (mm)
    "body_w": 29,      # AC W (mm)
    "body_h": 30,      # AD Max Height (mm)
}
FIRST_DATA_ROW = 11

DIM_KEYS = ("p2", "w1", "w2", "l", "h_max", "pcb_type", "source", "status", "note")


def to_num(value):
    """Turn a cell value into a float, or None. Accepts '6,4' and '6.4'."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def to_text(value):
    if value is None:
        return ""
    return str(value).strip()


def load_rows(path, key_field):
    rows = {}
    if not path.exists():
        return rows
    with path.open(newline="", encoding="utf-8") as handle:
        for record in csv.DictReader(row for row in handle if not row.lstrip().startswith("#")):
            key = to_text(record.get(key_field)).upper()
            if not key:
                continue
            entry = {}
            for name in DIM_KEYS:
                raw = to_text(record.get(name))
                if name in ("pcb_type", "source", "status", "note"):
                    entry[name] = raw
                else:
                    entry[name] = to_num(raw)
            rows[key] = entry
    return rows


def load_parts():
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - guided failure
        sys.exit(
            "openpyxl is required.\n"
            "  python3 -m venv .venv && .venv/bin/pip install openpyxl\n"
            "  .venv/bin/python tools/build_catalog.py"
        )

    workbook = openpyxl.load_workbook(XLSX, data_only=True)
    sheet = workbook[SHEET]
    parts = []
    for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
        pn = to_text(sheet.cell(row, COL["pn"]).value)
        if not pn:
            continue
        if not to_text(sheet.cell(row, COL["catalog"]).value):
            continue  # not a catalogue (web published) part
        parts.append(
            {
                "pn": pn.upper(),
                "suffix": pn[-3:].upper(),
                "series": to_text(sheet.cell(row, COL["series"]).value),
                "size": to_text(sheet.cell(row, COL["size"]).value),
                "feature": to_text(sheet.cell(row, COL["feature"]).value),
                "state": to_text(sheet.cell(row, COL["status"]).value),
                "body_l": to_num(sheet.cell(row, COL["body_l"]).value),
                "body_w": to_num(sheet.cell(row, COL["body_w"]).value),
                "body_h": to_num(sheet.cell(row, COL["body_h"]).value),
            }
        )
    parts.sort(key=lambda item: item["pn"])
    return parts


def build_catalog():
    suffixes = load_rows(DIMS_CSV, "suffix")
    overrides = load_rows(OVERRIDES_CSV, "pn")
    parts = load_parts()

    for part in parts:
        dims = overrides.get(part["pn"]) or suffixes.get(part["suffix"])
        part["dims"] = dims or None
        part["dims_basis"] = "pn-override" if part["pn"] in overrides else ("suffix" if dims else "none")

    used = {part["suffix"] for part in parts}
    return {
        "generated": date.today().isoformat(),
        "source": {"workbook": XLSX.name, "sheet": SHEET},
        "suffixes": {key: value for key, value in sorted(suffixes.items()) if key in used},
        "parts": parts,
    }


def js_block(catalog):
    payload = json.dumps(catalog, ensure_ascii=False, separators=(",", ":"))
    return f"{BEGIN}\nwindow.PCC_CATALOG = {payload};\n{END}"


def inject(catalog, check_only=False):
    html = HTML.read_text(encoding="utf-8")
    if BEGIN not in html or END not in html:
        sys.exit(f"Markers missing in {HTML.name}. Expected {BEGIN} ... {END}")
    new_html = re.sub(
        re.escape(BEGIN) + r".*?" + re.escape(END),
        lambda _match: js_block(catalog),
        html,
        flags=re.S,
    )
    if check_only:
        if new_html != html:
            sys.exit(f"{HTML.name} is out of date. Re-run tools/build_catalog.py")
        print(f"{HTML.name} is up to date.")
        return
    if new_html != html:
        HTML.write_text(new_html, encoding="utf-8")
        print(f"Updated {HTML.name}")
    else:
        print(f"{HTML.name} unchanged")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--check", action="store_true", help="verify index.html is up to date, do not write")
    args = parser.parse_args()

    catalog = build_catalog()
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    parts = catalog["parts"]
    with_dims = [part for part in parts if part["dims"]]
    resolved = [part for part in parts if part["dims"] and part["dims"].get("w1") is not None]
    manual_suffixes = sorted(
        {part["suffix"] for part in parts if not part["dims"] or part["dims"].get("w1") is None}
    )

    print(f"Part numbers        : {len(parts)}")
    print(f"With a dimension row: {len(with_dims)}")
    print(f"With W1/W2/L values : {len(resolved)}")
    print(f"Manual suffixes     : {', '.join(manual_suffixes) if manual_suffixes else '-'}")
    print(f"Wrote {OUT_JSON.relative_to(ROOT)}")
    inject(catalog, check_only=args.check)


if __name__ == "__main__":
    main()
